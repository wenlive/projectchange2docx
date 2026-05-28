#!/usr/bin/env python3
"""
Regression verification script for git-diff-review Word documents.

Cross-references a generated .docx against the source git diff to verify:
  - Files are correctly classified (included / excluded by rule)
  - Path rewriting and content sanitization are applied
  - No brand terms leak through to the output

Usage:
    python3 verify_change_review.py <commit> <document.docx>
    python3 verify_change_review.py <commit> <document.docx> --sample N

Must be invoked from the target repository root (same as generate_change_review.py).
"""

import subprocess
import os
import sys
import fnmatch
import re
import random
import argparse
from collections import Counter, defaultdict

try:
    from docx import Document
except ImportError:
    print("Error: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

# ── Configuration (mirrors generate_change_review.py) ─────────────────────────

EXCLUDE_PATTERNS = [
    "node_modules", "__pycache__", ".git", ".svn", ".hg", ".cache",
    "*.pyc", "*.pyo", "*.so", "*.o", "*.a", "*.dll", "*.exe",
    "*.bin", "*.zip", "*.tar", "*.gz", "*.bz2", "*.7z", "*.rar",
    "*.png", "*.jpg", "*.jpeg", "*.gif", "*.ico", "*.svg", "*.pdf",
    "*.ttf", "*.woff", "*.woff2", "*.eot", "*.mp3", "*.mp4", "*.avi",
    "*.mov", "*.wasm", "*.class", "*.jar",
]

FORCE_INCLUDE_PREFIXES = []

SKIP_DOC_FILES = True
DOC_EXCLUDE_PATTERNS = [
    "*.md", "*.markdown", "*.MD",
    ".gitignore", "Third_Party_Open_Source_Software_Notice",
]

FILEPATH_EXCLUDE_KEYWORDS = [
    "opengauss", "gaussdb", "openeuler", "kunpeng",
    "mogdb", "vastbase", "huawei", "huaweicloud",
]

PATH_REWRITE_RULES = [
    ("gausskernel", "kernel"),
    ("GAUSSKERNEL", "KERNEL"),
]

CONTENT_REPLACEMENTS = [
    ("gausskernel", "kernel"),
    ("GAUSSKERNEL", "KERNEL"),
    ("OpenGauss", "HelmDB"),
    ("openGauss", "HelmDB"),
    ("OPENGAUSS", "HELMDB"),
    ("opengauss", "helmdb"),
    ("GaussDB", "HelmDB"),
    ("GAUSSDB", "HELMDB"),
    ("gaussdb", "helmdb"),
]

MIN_CHANGED_LINES = 20
FULL_OUTPUT_LINES = 200

# ── Helpers ──────────────────────────────────────────────────────────────────

def run(cmd_args, cwd=None):
    if cwd is None:
        cwd = os.getcwd()
    if cmd_args[0] == "git":
        cmd_args = ["git", "-c", "core.quotePath=false"] + cmd_args[1:]
    result = subprocess.run(cmd_args, capture_output=True, text=True, cwd=cwd)
    if result.returncode != 0:
        raise RuntimeError(f"Command failed: {' '.join(cmd_args)}\n{result.stderr}")
    return result.stdout


def is_excluded(filepath):
    if is_force_include(filepath):
        return False
    filepath = filepath.replace("\\", "/")
    segments = filepath.split("/")
    for seg in segments:
        for pat in EXCLUDE_PATTERNS:
            if fnmatch.fnmatch(seg, pat):
                return True
    for pat in EXCLUDE_PATTERNS:
        if "/" in pat and fnmatch.fnmatch(filepath, pat):
            return True
    return False


def is_force_include(filepath):
    filepath = filepath.replace("\\", "/")
    for prefix in FORCE_INCLUDE_PREFIXES:
        if filepath.startswith(prefix):
            return True
    return False


def is_doc_file(filepath):
    if not SKIP_DOC_FILES:
        return False
    if is_force_include(filepath):
        return False
    filepath = filepath.replace("\\", "/")
    filename = filepath.split("/")[-1] if "/" in filepath else filepath
    for pat in DOC_EXCLUDE_PATTERNS:
        if fnmatch.fnmatch(filename, pat):
            return True
    return False


def is_brand_excluded(filepath):
    if is_force_include(filepath):
        return False
    filepath = filepath.replace("\\", "/")
    segments = filepath.split("/")
    for seg in segments:
        seg_lower = seg.lower()
        for kw in FILEPATH_EXCLUDE_KEYWORDS:
            if kw in seg_lower:
                return True
    return False


def rewrite_filepath(filepath):
    result = filepath
    for old, new in PATH_REWRITE_RULES:
        result = result.replace(old, new)
    return result


def sanitize_content(text):
    result = text
    for old, new in CONTENT_REPLACEMENTS:
        result = result.replace(old, new)
    return result


def clean_display_path(filepath):
    return sanitize_content(rewrite_filepath(filepath))


def normalize_content(text):
    """Normalize content for comparison: the Word doc uses a single space
    as placeholder for empty lines. Mirror this in git-side content."""
    lines = text.split("\n")
    return "\n".join(" " if line == "" else line for line in lines)


def parse_numstat(output):
    stats = {}
    for line in output.strip().split("\n"):
        if not line:
            continue
        parts = line.split("\t")
        if len(parts) < 3:
            continue
        added, deleted = parts[0], parts[1]
        raw_path = "\t".join(parts[2:])

        old_path = None
        m = re.match(r'^(.*)\{(.*?) => (.*?)\}(.*)$', raw_path)
        if m:
            prefix, old_part, new_part, suffix = m.group(1), m.group(2).strip(), m.group(3).strip(), m.group(4)
            filepath = (prefix + new_part + suffix).replace('//', '/')
            if old_part:
                old_path = (prefix + old_part + suffix).replace('//', '/')
        else:
            filepath = raw_path

        if added == "-" and deleted == "-":
            stats[filepath] = {"added": 0, "deleted": 0, "binary": True, "old_path": old_path}
        else:
            stats[filepath] = {
                "added": int(added) if added.lstrip("-").isdigit() else 0,
                "deleted": int(deleted) if deleted.lstrip("-").isdigit() else 0,
                "binary": False,
                "old_path": old_path,
            }
    return stats


def parse_name_status(output):
    statuses = {}
    for line in output.strip().split("\n"):
        if not line:
            continue
        parts = line.split("\t")
        code = parts[0]
        if code.startswith("R"):
            old_p, new_p = parts[1], parts[2]
            statuses[new_p] = {"status": "R", "old_path": old_p}
        elif code.startswith("C"):
            old_p, new_p = parts[1], parts[2]
            statuses[new_p] = {"status": "C", "old_path": old_p}
        else:
            filepath = parts[1]
            statuses[filepath] = {"status": code, "old_path": None}
    return statuses


def read_head_file(filepath):
    try:
        out = run(["git", "show", f"HEAD:{filepath}"])
        return out, None
    except Exception as e:
        return None, str(e)


# ── Word document parsing ────────────────────────────────────────────────────

def parse_docx_headings(docx_path):
    """Extract file entries from Word document headings.
    Returns dict: display_path → heading_text
    """
    doc = Document(docx_path)
    entries = {}
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            text = para.text.strip()
            # File headings look like: "path/to/file  [STATUS]  (+N/-M)"
            # Appendix headings are separate
            if text.startswith("Appendix") or "项目代码整合文档" in text:
                continue
            # Extract the path part (path may contain spaces; terminated by 2+ spaces then "[")
            match = re.match(r'^(.+?)\s{2,}\[', text)
            if match:
                entries[match.group(1)] = text
    return entries


def parse_docx_appendix(docx_path, appendix_title):
    """Extract file paths from a named appendix section.
    Returns list of file paths.
    """
    doc = Document(docx_path)
    in_appendix = False
    paths = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if para.style.name.startswith('Heading') and appendix_title in text:
            in_appendix = True
            continue
        if para.style.name.startswith('Heading') and in_appendix:
            break
        if in_appendix and text and not text.startswith("The following"):
            # Remove leading spaces
            path = text.lstrip()
            if path:
                paths.append(path)
    return paths


def extract_content_sample(docx_path, sample_paths):
    """Extract content blocks for given display paths from the Word doc.
    Returns dict: display_path → content_text
    """
    doc = Document(docx_path)
    samples = {}
    current_file = None
    code_paras = []  # collect run texts from all code paragraphs for current file

    for para in doc.paragraphs:
        text = para.text.strip()
        if para.style.name.startswith('Heading'):
            # Save previous file's content
            if current_file and current_file in sample_paths and code_paras:
                samples[current_file] = "\n".join(code_paras)
            current_file = None
            code_paras = []
            # Check if this is a file heading
            match = re.match(r'^(\S+)\s+\[', text)
            if match:
                current_file = match.group(1)
            continue

        if current_file and current_file in sample_paths:
            if para.runs and para.runs[0].font.name == "Consolas":
                for run in para.runs:
                    # python-docx renders <w:br/> as \n appended to the run text.
                    # Strip trailing \n for comparison purposes.
                    line = run.text.rstrip("\n")
                    code_paras.append(line)

    # Don't forget the last file
    if current_file and current_file in sample_paths and code_paras:
        samples[current_file] = "\n".join(code_paras)

    return samples


# ── Brand term scanning ──────────────────────────────────────────────────────

BRAND_TERMS = [
    "gausskernel", "GAUSSKERNEL",
    "openGauss", "OpenGauss", "OPENGAUSS", "opengauss",
    "GaussDB", "GAUSSDB", "gaussdb",
]


def scan_docx_for_brand_terms(docx_path):
    """Scan entire Word document for remaining brand terms.
    Returns dict: term → [(para_index, snippet)]
    """
    doc = Document(docx_path)
    findings = defaultdict(list)
    for i, para in enumerate(doc.paragraphs):
        text = para.text
        for term in BRAND_TERMS:
            if term in text:
                # Find context around the term
                idx = text.find(term)
                start = max(0, idx - 30)
                end = min(len(text), idx + len(term) + 30)
                snippet = ("..." if start > 0 else "") + text[start:end] + ("..." if end < len(text) else "")
                findings[term].append((i, para.style.name, snippet))
    return dict(findings)


# ── Main verification ────────────────────────────────────────────────────────

def classify_from_git(commit):
    """Classify all files from git diff using the same rules as generate_change_review.py.
    Returns dicts of categorized file lists.
    """
    statuses_raw = run(["git", "diff", "--name-status", f"{commit}^..HEAD"])
    statuses = parse_name_status(statuses_raw)

    numstats_raw = run(["git", "diff", "--numstat", f"{commit}^..HEAD"])
    numstats = parse_numstat(numstats_raw)

    included = {}       # filepath → info (should be in Word doc)
    brand_excluded = [] # filepath (should be in brand appendix)
    doc_excluded = []   # filepath (should be in doc appendix)
    binary_excluded = []# filepath (should be in binary appendix)
    excluded_other = [] # filepath (pattern or too-few-lines)
    deleted = []        # filepath

    for filepath, ns in sorted(numstats.items()):
        # Pattern exclude
        if is_excluded(filepath):
            excluded_other.append(filepath)
            continue

        # Doc exclude
        if is_doc_file(filepath):
            doc_excluded.append(filepath)
            continue

        # Brand exclude
        if is_brand_excluded(filepath):
            brand_excluded.append(filepath)
            continue

        # Deleted
        st = statuses.get(filepath, {})
        if st.get("status") == "D":
            deleted.append(filepath)
            continue

        # Binary
        if ns["binary"]:
            binary_excluded.append(filepath)
            continue

        # Line count threshold
        changed = ns["added"] + ns["deleted"]
        force = is_force_include(filepath)
        if not force and changed < MIN_CHANGED_LINES:
            excluded_other.append(filepath)
            continue

        st = statuses.get(filepath, {"status": "M", "old_path": None})
        included[filepath] = {
            **ns,
            "status": st["status"],
            "changed_lines": changed,
        }

    return {
        "included": included,
        "brand_excluded": brand_excluded,
        "doc_excluded": doc_excluded,
        "binary_excluded": binary_excluded,
        "excluded_other": excluded_other,
        "deleted": deleted,
        "total": len(numstats),
    }


def verify(commit, docx_path, sample_size=5):
    """Main verification routine."""
    print("=" * 72)
    print("  git-diff-review Regression Verification")
    print(f"  Commit : {commit}")
    print(f"  Docx   : {docx_path}")
    print("=" * 72)

    # 1. Classify from git
    print("\n[1/4] Classifying files from git diff ...", flush=True)
    git_class = classify_from_git(commit)
    n_included = len(git_class["included"])
    n_brand = len(git_class["brand_excluded"])
    n_doc = len(git_class["doc_excluded"])
    n_binary = len(git_class["binary_excluded"])
    n_other = len(git_class["excluded_other"])
    n_deleted = len(git_class["deleted"])

    print(f"  Total files in diff  : {git_class['total']}")
    print(f"  Expected in doc      : {n_included}")
    print(f"  Brand excluded       : {n_brand}")
    print(f"  Doc excluded         : {n_doc}")
    print(f"  Binary excluded      : {n_binary}")
    print(f"  Other excluded       : {n_other}  (<{MIN_CHANGED_LINES} lines or pattern)")
    print(f"  Deleted              : {n_deleted}")
    print(f"  Sum                  : {n_included + n_brand + n_doc + n_binary + n_other + n_deleted}")

    # 2. Parse Word document
    print("\n[2/4] Parsing Word document ...", flush=True)
    doc_headings = parse_docx_headings(docx_path)
    doc_brand_appendix = parse_docx_appendix(docx_path, "Brand-Excluded")
    doc_doc_appendix = parse_docx_appendix(docx_path, "Skipped Document")
    doc_binary_appendix = parse_docx_appendix(docx_path, "Skipped Binary")

    print(f"  File headings in doc : {len(doc_headings)}")
    print(f"  Brand appendix paths : {len(doc_brand_appendix)}")
    print(f"  Doc appendix paths   : {len(doc_doc_appendix)}")
    print(f"  Binary appendix paths: {len(doc_binary_appendix)}")

    # 3. Cross-reference
    print("\n[3/4] Cross-referencing ...", flush=True)
    anomalies = []

    # --- 3a. Expected included files ---
    expected_display_paths = set()
    for fp in git_class["included"]:
        expected_display_paths.add(clean_display_path(fp))

    actual_display_paths = set(doc_headings.keys())

    missing_from_doc = expected_display_paths - actual_display_paths
    extra_in_doc = actual_display_paths - expected_display_paths

    if missing_from_doc:
        anomalies.append(f"  MISSING from doc ({len(missing_from_doc)} files):")
        for fp in sorted(missing_from_doc)[:10]:
            anomalies.append(f"    - {fp}")
        if len(missing_from_doc) > 10:
            anomalies.append(f"    ... and {len(missing_from_doc) - 10} more")

    if extra_in_doc:
        anomalies.append(f"  EXTRA in doc ({len(extra_in_doc)} files):")
        for fp in sorted(extra_in_doc)[:10]:
            anomalies.append(f"    + {fp}")
        if len(extra_in_doc) > 10:
            anomalies.append(f"    ... and {len(extra_in_doc) - 10} more")

    # --- 3b. Brand-excluded files ---
    expected_brand_display = {clean_display_path(fp) for fp in git_class["brand_excluded"]}
    actual_brand_display = set(doc_brand_appendix)

    # Brand files that should be listed but aren't
    brand_missing_from_appendix = expected_brand_display - actual_brand_display
    # Brand files listed in appendix that shouldn't be
    brand_extra_in_appendix = actual_brand_display - expected_brand_display

    # Also check that brand-excluded files do NOT appear in doc headings
    brand_leaked_to_doc = expected_brand_display & actual_display_paths
    if brand_leaked_to_doc:
        anomalies.append(f"  BRAND LEAK ({len(brand_leaked_to_doc)} brand-excluded files appeared in doc!):")
        for fp in sorted(brand_leaked_to_doc)[:5]:
            anomalies.append(f"    !! {fp}")

    # --- 3c. Content sanitization: sample check ---
    print(f"\n[4/4] Sampling {sample_size} files for content verification ...", flush=True)

    # Pick sample files: prefer files with original paths containing gausskernel
    gk_files = [fp for fp in git_class["included"] if "gausskernel" in fp.lower()]
    other_files = [fp for fp in git_class["included"] if "gausskernel" not in fp.lower()]
    sample_sources = gk_files[:sample_size] + other_files[:max(0, sample_size - len(gk_files))]
    if len(sample_sources) > sample_size:
        sample_sources = random.sample(sample_sources, sample_size)

    # Build mapping: display_path → original_path
    display_to_original = {}
    for fp in git_class["included"]:
        display_to_original[clean_display_path(fp)] = fp

    sample_display_paths = [clean_display_path(fp) for fp in sample_sources]
    doc_samples = extract_content_sample(docx_path, set(sample_display_paths))

    content_issues = []
    content_ok = 0

    for fp, display_p in zip(sample_sources, sample_display_paths):
        git_content, err = read_head_file(fp)
        if err:
            content_issues.append(f"    Cannot read {fp}: {err}")
            continue

        expected_content = sanitize_content(git_content)
        doc_content = doc_samples.get(display_p)

        if doc_content is None:
            content_issues.append(f"    Content NOT FOUND in doc for: {display_p}")
            continue

        # Normalize both sides for empty-line placeholder comparison
        if normalize_content(expected_content) != normalize_content(doc_content):
            # Find first differing line
            exp_lines = expected_content.split("\n")
            doc_lines = doc_content.split("\n")
            diff_at = None
            for i, (el, dl) in enumerate(zip(exp_lines, doc_lines)):
                if el != dl:
                    diff_at = i + 1
                    break
            if diff_at is None and len(exp_lines) != len(doc_lines):
                diff_at = f"length mismatch (expected {len(exp_lines)}, got {len(doc_lines)})"
            content_issues.append(
                f"    Content MISMATCH: {display_p}\n"
                f"      First diff at line: {diff_at}"
            )
        else:
            content_ok += 1

    # --- 3d. Brand term scan ---
    brand_findings = scan_docx_for_brand_terms(docx_path)

    # 4. Print report
    print("\n" + "=" * 72)
    print("  VERIFICATION REPORT")
    print("=" * 72)

    print(f"\n  Included files : {n_included:>6} expected, {len(doc_headings):>6} in doc  "
          f"{'✓' if len(missing_from_doc) == 0 and len(extra_in_doc) == 0 else '✗'}")

    print(f"  Brand excluded : {n_brand:>6} expected, {len(doc_brand_appendix):>6} in appendix  "
          f"{'✓' if len(brand_leaked_to_doc) == 0 else '✗'}")

    print(f"  Doc excluded   : {n_doc:>6} expected, {len(doc_doc_appendix):>6} in appendix")
    print(f"  Binary excluded: {n_binary:>6} expected, {len(doc_binary_appendix):>6} in appendix")
    print(f"  Other excluded : {n_other:>6}  (not in doc)")
    print(f"  Deleted files  : {n_deleted:>6}  (not in doc)")

    # Content sampling
    print(f"\n  Content samples verified: {content_ok}/{len(sample_sources)} matched")
    if content_issues:
        print(f"  Content issues ({len(content_issues)}):")
        for issue in content_issues:
            print(issue)

    # Brand term scan
    print(f"\n  Brand term scan:")
    if brand_findings:
        for term, occurrences in sorted(brand_findings.items()):
            print(f"    '{term}' found in {len(occurrences)} paragraph(s):")
            for para_idx, style, snippet in occurrences[:5]:
                print(f"      Para {para_idx} [{style}]: {snippet}")
            if len(occurrences) > 5:
                print(f"      ... and {len(occurrences) - 5} more")
    else:
        print(f"    ✓ No brand terms found — document is clean!")

    # Anomalies
    if anomalies:
        print(f"\n  Anomalies ({len(anomalies)}):")
        for a in anomalies:
            print(a)
    else:
        print(f"\n  ✓ No anomalies detected!")

    # Summary verdict
    print(f"\n  ---")
    total_issues = (
        len(missing_from_doc) + len(extra_in_doc) +
        len(brand_leaked_to_doc) + len(content_issues) +
        len(brand_findings)
    )
    if total_issues == 0:
        print(f"  VERDICT: PASS — document is consistent with git diff rules.")
    else:
        print(f"  VERDICT: REVIEW — {total_issues} potential issue(s) found.")

    print("=" * 72)


# ── CLI ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Verify a git-diff-review Word document against source git diff."
    )
    parser.add_argument("commit", help="Target commit hash (same as used for generate_change_review.py)")
    parser.add_argument("docx", help="Path to the generated .docx file")
    parser.add_argument("--sample", type=int, default=5,
                        help="Number of files to sample for content verification (default: 5)")
    args = parser.parse_args()

    if not os.path.isfile(args.docx):
        print(f"Error: document not found: {args.docx}")
        sys.exit(1)

    verify(args.commit, args.docx, args.sample)


if __name__ == "__main__":
    main()
